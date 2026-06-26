"""生成 MinIO 测试数据 — PDF/Word 文档 + 上传到 MinIO"""
import os
import sys
import io

def generate_test_pdf():
    """生成小型测试 PDF (~1页, ~100字)"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        print("⚠ reportlab 未安装，用 PyMuPDF 生成替代 PDF")
        # 用 fitz 创建简单 PDF
        try:
            import fitz
            doc = fitz.open()
            page = doc.new_page()
            text = "OpenClaw E2E 测试文档\n\n这是一份用于端到端测试的 PDF 文档。\n" \
                   "文档内容包括：系统架构说明、部署配置、性能指标。\n" \
                   "核心服务：Proxy Pod、LiteLLM Proxy、Stream Service。\n" \
                   "支持的附件类型：PDF、Word、音频、图片。\n" \
                   "MinIO 用于文件存储和预签名 URL 生成。"
            page.insert_text((72, 72), text, fontsize=12)
            pdf_bytes = doc.tobytes()
            doc.close()
            return pdf_bytes
        except ImportError:
            print("❌ fitz 也未安装，生成空 PDF")
            return b""

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    text = "OpenClaw E2E 测试文档\n\n这是一份用于端到端测试的 PDF 文档。" \
           "\n文档内容包括：系统架构说明、部署配置、性能指标。" \
           "\n核心服务：Proxy Pod、LiteLLM Proxy、Stream Service。" \
           "\n支持的附件类型：PDF、Word、音频、图片。" \
           "\nMinIO 用于文件存储和预签名 URL 生成。"
    c.setFont("Helvetica", 12)
    y = A4[1] - 72
    for line in text.split("\n"):
        c.drawString(72, y, line)
        y -= 20
    c.showPage()
    c.save()
    return buf.getvalue()


def generate_test_docx():
    """生成小型测试 Word 文档 (~1段落, ~100字)"""
    try:
        from docx import Document
    except ImportError:
        print("❌ python-docx 未安装，无法生成 Word 文档")
        return b""

    doc = Document()
    doc.add_heading('OpenClaw E2E 测试文档', level=1)
    doc.add_paragraph(
        '这是一份用于端到端测试的 Word 文档。'
        '文档内容包括：系统架构说明、部署配置、性能指标。'
        '核心服务：Proxy Pod、LiteLLM Proxy、Stream Service。'
        '支持的附件类型：PDF、Word、音频、图片。'
        'MinIO 用于文件存储和预签名 URL 生成。'
    )
    doc.add_paragraph('测试环境：K8S KinD 集群，单 GPU，Ollama 本地推理。')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_to_minio(bucket, files_dict):
    """上传文件到 MinIO（通过 mc 命令行或 Python minio SDK）"""
    endpoint = os.getenv("MINIO_ENDPOINT", "localhost:9000")
    access_key = os.getenv("MINIO_ACCESS_KEY", "minioadmin")
    secret_key = os.getenv("MINIO_SECRET_KEY", "minioadmin")

    try:
        from minio import Minio
        client = Minio(endpoint, access_key=access_key, secret_key=secret_key, secure=False)

        # 确保 bucket 存在
        if not client.bucket_exists(bucket):
            client.make_bucket(bucket)
            print(f"✓ Created bucket: {bucket}")

        for key, data in files_dict.items():
            client.put_object(
                bucket, key, io.BytesIO(data), len(data),
                content_type=_get_content_type(key)
            )
            print(f"✓ Uploaded: {bucket}/{key} ({len(data)} bytes)")

        return True
    except ImportError:
        print("⚠ minio SDK 未安装，尝试 mc 命令行")
        return _upload_via_mc(bucket, files_dict)
    except Exception as e:
        print(f"❌ MinIO 上传失败: {e}")
        return False


def _get_content_type(key):
    if key.endswith('.pdf'):
        return 'application/pdf'
    elif key.endswith('.docx'):
        return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif key.endswith('.wav'):
        return 'audio/wav'
    elif key.endswith('.png'):
        return 'image/png'
    return 'application/octet-stream'


def _upload_via_mc(bucket, files_dict):
    """通过 mc 命令行工具上传（备选方案）"""
    import tempfile
    import subprocess

    endpoint = os.getenv("MINIO_ENDPOINT", "localhost:9000")
    access_key = os.getenv("MINIO_ACCESS_KEY", "minioadmin")
    secret_key = os.getenv("MINIO_SECRET_KEY", "minioadmin")

    # 设置 mc alias
    try:
        subprocess.run(["mc", "alias", "set", "local",
                        f"http://{endpoint}", access_key, secret_key],
                       capture_output=True, timeout=10)
    except FileNotFoundError:
        print("❌ mc 命令行工具未安装")
        return False

    # 确保 bucket 存在
    subprocess.run(["mc", "mb", f"local/{bucket}", "--ignore-existing"],
                   capture_output=True, timeout=10)

    for key, data in files_dict.items():
        with tempfile.NamedTemporaryFile(suffix=os.path.splitext(key)[1], delete=False) as f:
            f.write(data)
            f.flush()
            tmp_path = f.name
        try:
            subprocess.run(["mc", "cp", tmp_path, f"local/{bucket}/{key}"],
                           capture_output=True, timeout=30)
            print(f"✓ Uploaded via mc: {bucket}/{key}")
        finally:
            os.unlink(tmp_path)

    return True


def main():
    print("── 生成测试数据 ──")

    pdf_bytes = generate_test_pdf()
    docx_bytes = generate_test_docx()

    print(f"  PDF: {len(pdf_bytes)} bytes")
    print(f"  DOCX: {len(docx_bytes)} bytes")

    files_dict = {}
    if pdf_bytes:
        files_dict['test_document.pdf'] = pdf_bytes
    if docx_bytes:
        files_dict['test_document.docx'] = docx_bytes

    if not files_dict:
        print("❌ 无法生成测试文件")
        return

    print("\n── 上传到 MinIO ──")
    success = upload_to_minio('openclaw-test', files_dict)
    if success:
        print("\n✓ 测试数据已上传到 MinIO openclaw-test bucket")
    else:
        print("\n❌ 上传失败，请手动上传或检查 MinIO 配置")
        # 保存到本地作为 fallback
        for key, data in files_dict.items():
            local_path = f"/tmp/{key}"
            with open(local_path, 'wb') as f:
                f.write(data)
            print(f"  已保存到: {local_path}")


if __name__ == "__main__":
    main()
